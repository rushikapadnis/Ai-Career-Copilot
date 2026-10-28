
from io import BytesIO

import pytest
from docx import Document

from app.services.resume_parser import (
    ResumeParsingError,
    extract_resume_text,
)


def create_test_docx() -> bytes:
    """Create an in-memory DOCX resume for testing."""
    document = Document()

    document.add_heading("Candidate Resume", level=1)
    document.add_paragraph("Name: Test Candidate")
    document.add_paragraph("Skills: Python, SQL, Machine Learning")
    document.add_paragraph("Experience: 2 years")

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def test_extract_docx_text():
    file_bytes = create_test_docx()

    text = extract_resume_text(
        "sample_resume.docx",
        file_bytes,
    )

    assert "Test Candidate" in text
    assert "Python" in text
    assert "Machine Learning" in text
    assert "2 years" in text


def test_reject_unsupported_file():
    with pytest.raises(ResumeParsingError):
        extract_resume_text(
            "resume.txt",
            b"Some resume text",
        )


def test_reject_empty_file():
    with pytest.raises(ResumeParsingError):
        extract_resume_text(
            "resume.pdf",
            b"",
        )


def test_reject_file_larger_than_limit():
    large_file = b"x" * (5 * 1024 * 1024 + 1)

    with pytest.raises(ResumeParsingError):
        extract_resume_text(
            "resume.pdf",
            large_file,
        )


def test_reject_docx_without_readable_text():
    document = Document()
    buffer = BytesIO()
    document.save(buffer)

    with pytest.raises(ResumeParsingError):
        extract_resume_text(
            "empty_resume.docx",
            buffer.getvalue(),
        )
