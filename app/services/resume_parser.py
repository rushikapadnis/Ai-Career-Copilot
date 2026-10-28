
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB


class ResumeParsingError(Exception):
    """Raised when a resume cannot be parsed safely."""


def extract_resume_text(
    file_name: str,
    file_bytes: bytes,
) -> str:
    """
    Extract text from a PDF or DOCX resume.

    Args:
        file_name: Original uploaded filename.
        file_bytes: Raw file content.

    Returns:
        Extracted and normalized resume text.

    Raises:
        ResumeParsingError: If the file is unsupported,
        empty, too large, or cannot be parsed.
    """
    extension = Path(file_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParsingError(
            "Unsupported file type. Upload a PDF or DOCX file."
        )

    if not file_bytes:
        raise ResumeParsingError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise ResumeParsingError(
            "Resume exceeds the 5 MB upload limit."
        )

    try:
        if extension == ".pdf":
            text = _extract_pdf_text(file_bytes)
        else:
            text = _extract_docx_text(file_bytes)

    except ResumeParsingError:
        raise
    except Exception as exc:
        raise ResumeParsingError(
            "Unable to read the resume. "
            "Please check that the file is valid."
        ) from exc

    # Normalize whitespace while preserving paragraph separation.
    lines = [
        " ".join(line.split())
        for line in text.splitlines()
    ]
    cleaned_text = "\n".join(
        line for line in lines if line
    ).strip()

    if not cleaned_text:
        raise ResumeParsingError(
            "No readable text was found in the resume. "
            "Scanned PDFs may require OCR."
        )

    return cleaned_text


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))

    if reader.is_encrypted:
        raise ResumeParsingError(
            "Password-protected PDFs are not supported."
        )

    pages = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        pages.append(page_text)

    return "\n".join(pages)


def _extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # Include text from tables, which some resumes use
    # for skills, education, or work experience.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)
